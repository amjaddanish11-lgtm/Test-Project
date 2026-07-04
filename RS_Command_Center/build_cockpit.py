#!/usr/bin/env python3
"""Build RS_Command_Center.xlsx from BUILD_SPEC.md, driven by rs_checks.json."""

import json
import sys
from datetime import datetime, date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Protection
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
DATA_FILE = HERE / "rs_checks.json"
OUT_FILE = HERE / "RS_Command_Center.xlsx"

# ---------------------------------------------------------------------------
# Palette / constants
# ---------------------------------------------------------------------------
NAVY = "FF1F2A44"
GOLD = "FFC9A227"
BEIGE = "FFE8E2D4"
CREAM = "FFF4F1EA"
BLUE_INPUT = "FF0000FF"
GRAY = "FF666666"
RED = "FFCC0000"
GREEN = "FF38761D"
AMBER = "FFB45309"
WHITE = "FFFFFFFF"
NAV_GRAY = "FF8A8578"
DARK_TEXT = "FF222222"
TILE_BORDER_COLOR = "FFD9D2C0"

SHEET_ORDER = ["Start Here", "Dashboard", "Pipeline", "Jobs", "Invoices", "Settings"]

MONEY_FMT = "#,##0.00"
PCT_FMT = "0.0%"
DATE_FMT = "DD/MM/YYYY"
COUNT_FMT = "0"

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_GOLD = PatternFill("solid", fgColor=GOLD)
FILL_BEIGE = PatternFill("solid", fgColor=BEIGE)
FILL_CREAM = PatternFill("solid", fgColor=CREAM)

NAV_LABELS = [
    ("B1", "Start", "Start Here"),
    ("D1", "Dashboard", "Dashboard"),
    ("F1", "Pipeline", "Pipeline"),
    ("H1", "Jobs", "Jobs"),
    ("J1", "Invoices", "Invoices"),
    ("L1", "Settings", "Settings"),
]


def parse_date(s):
    if s is None:
        return None
    return datetime.strptime(s, "%Y-%m-%d").date()


def load_data():
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        payload = json.load(f)
    return payload["dataset"]


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def apply_common(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False


def build_nav_row(ws, current_sheet_name):
    ws.row_dimensions[1].height = 15
    for coord, label, target_sheet in NAV_LABELS:
        cell = ws[coord]
        cell.value = label
        if target_sheet == current_sheet_name:
            cell.font = Font(name="Arial", size=8, bold=True, color=NAVY)
        else:
            cell.font = Font(name="Arial", size=8, color=NAV_GRAY)
        cell.hyperlink = Hyperlink(ref=coord, location="'{}'!A1".format(target_sheet))


def build_banner(ws, merge_range, text, size=14):
    ws.merge_cells(merge_range)
    top_left = merge_range.split(":")[0]
    cell = ws[top_left]
    cell.value = text
    cell.font = Font(name="Arial", size=size, bold=True, color=WHITE)
    cell.fill = FILL_NAVY
    for row in ws[merge_range]:
        for c in row:
            c.fill = FILL_NAVY
    ws.row_dimensions[2].height = 24


def build_ribbon(ws, merge_range, text):
    ws.merge_cells(merge_range)
    top_left = merge_range.split(":")[0]
    cell = ws[top_left]
    cell.value = text
    cell.font = Font(name="Arial", size=9, bold=True, color=NAVY)
    cell.fill = FILL_GOLD
    for row in ws[merge_range]:
        for c in row:
            c.fill = FILL_GOLD
    ws.row_dimensions[3].height = 16


def input_style(cell, number_format=None):
    cell.fill = FILL_CREAM
    cell.font = Font(name="Arial", size=10, color=BLUE_INPUT)
    cell.protection = Protection(locked=False)
    if number_format:
        cell.number_format = number_format


def header_cell(cell, text, size=9):
    cell.value = text
    cell.fill = FILL_BEIGE
    cell.font = Font(name="Arial", size=size, bold=True, color=NAVY)


# ---------------------------------------------------------------------------
# Settings sheet
# ---------------------------------------------------------------------------

def build_settings(wb, data):
    ws = wb.create_sheet("Settings")
    apply_common(ws)
    build_nav_row(ws, "Settings")
    build_banner(ws, "B2:K2", "  SETTINGS")
    build_ribbon(ws, "B3:K3", "  Rename lists here — dropdowns and formulas follow automatically.")

    settings = data["settings"]

    ws["B5"].value = "Currency label"
    ws["B5"].font = Font(name="Arial", size=9, color=GRAY)
    input_style(ws["C5"])
    ws["C5"].value = "AED"

    ws["B6"].value = "VAT note"
    ws["B6"].font = Font(name="Arial", size=9, color=GRAY)
    input_style(ws["C6"])
    ws["C6"].value = "VAT 5% — prices exclude VAT unless stated"

    header_cell(ws["C9"], "Business lines", size=10)
    for i, line in enumerate(settings["business_lines"]):
        cell = ws.cell(row=10 + i, column=3)
        input_style(cell)
        cell.value = line

    header_cell(ws["E9"], "Lead sources", size=10)
    for i, src in enumerate(settings["sources"]):
        cell = ws.cell(row=10 + i, column=5)
        input_style(cell)
        cell.value = src

    header_cell(ws["G9"], "Stage", size=10)
    header_cell(ws["H9"], "Probability", size=10)
    for i, st in enumerate(settings["stages"]):
        gcell = ws.cell(row=10 + i, column=7)
        input_style(gcell)
        gcell.value = st["stage"]
        hcell = ws.cell(row=10 + i, column=8)
        input_style(hcell, number_format=PCT_FMT)
        hcell.value = st["probability"]

    header_cell(ws["J9"], "Job status", size=10)
    for i, status in enumerate(["In progress", "Done"]):
        cell = ws.cell(row=10 + i, column=10)
        input_style(cell)
        cell.value = status

    ws["B17"].value = "Stages Won and Lost count as closed. A lead in any other stage is open pipeline."
    ws["B17"].font = Font(name="Arial", size=8, color=GRAY)

    widths = {"B": 18, "C": 22, "D": 3, "E": 22, "F": 3, "G": 16, "H": 12, "I": 3, "J": 14, "K": 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    return ws


# ---------------------------------------------------------------------------
# Start Here sheet
# ---------------------------------------------------------------------------

def build_start_here(wb):
    ws = wb.create_sheet("Start Here")
    apply_common(ws)
    build_nav_row(ws, "Start Here")
    build_banner(ws, "B2:K2", "  RELIABLE SOURCES — COMMAND CENTER", size=16)
    build_ribbon(ws, "B3:K3", "  Who to chase today, what work is live, who owes you money — in under a minute.")

    def line(row, text, bold=False, color=DARK_TEXT, size=10):
        cell = ws.cell(row=row, column=2)
        cell.value = text
        cell.font = Font(name="Arial", size=size, bold=bold, color=color)

    line(5, "The 5-Minute Morning Routine", bold=True, color=NAVY, size=12)
    line(7, "1.  Open the Dashboard tab and read the tiles — they update themselves.")
    line(8, "2.  Check 'Follow-ups due' — call or WhatsApp each CALL TODAY lead in the Pipeline tab, update its stage, and set the next follow-up date.")
    line(9, "3.  Glance at overdue invoices on the Dashboard — chase the oldest one first.")
    line(10, "4.  New enquiry since yesterday? Log it in the Pipeline tab — one row, 30 seconds.")
    line(12, "What each tab does", bold=True, color=NAVY, size=12)
    line(13, "Dashboard — your morning read. All formulas, nothing to type.")
    line(14, "Pipeline — every enquiry and quote. Blue cells are yours to type; grey ones fill themselves.")
    line(15, "Jobs — won work only, with margin worked out for you.")
    line(16, "Invoices — what you've billed, what's been paid, what's overdue.")
    line(17, "Settings — rename business lines, sources and stages here without breaking anything.")
    line(19, "The one rule: a lead without a next follow-up date is a lead you've decided to lose.", bold=True, color=RED, size=11)
    line(21, "Blue writing on cream = type here. Everything else is locked so it can't break. Dummy data is loaded — replace it with your real pipeline.", color=GRAY, size=8)

    ws.column_dimensions["B"].width = 110

    return ws


# ---------------------------------------------------------------------------
# Pipeline sheet
# ---------------------------------------------------------------------------

def build_pipeline(wb, data):
    ws = wb.create_sheet("Pipeline")
    apply_common(ws)
    build_nav_row(ws, "Pipeline")
    build_banner(ws, "B2:N2", "  PIPELINE — every enquiry lives here")
    build_ribbon(ws, "B3:N3", "  Blue cells are yours. Probability, weighted value and the CALL TODAY flag fill themselves.")

    headers = ["Date in", "Client", "Contact", "Business line", "Source", "What they want",
               "Est. value AED", "Stage", "Probability %", "Weighted AED", "Next follow-up",
               "Status flag", "Notes"]
    thin_navy = Side(style="thin", color=NAVY)
    for i, h in enumerate(headers):
        cell = ws.cell(row=4, column=2 + i)
        header_cell(cell, h)
        cell.border = Border(bottom=thin_navy)

    rows = data["pipeline"]
    for i, rec in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=parse_date(rec["date_in"]))
        ws.cell(row=r, column=3, value=rec["client"])
        ws.cell(row=r, column=4, value=rec["contact"])
        ws.cell(row=r, column=5, value=rec["line"])
        ws.cell(row=r, column=6, value=rec["source"])
        ws.cell(row=r, column=7, value=rec["want"])
        ws.cell(row=r, column=8, value=rec["value"])
        ws.cell(row=r, column=9, value=rec["stage"])
        ws.cell(row=r, column=12, value=parse_date(rec["next_followup"]))
        ws.cell(row=r, column=14, value=rec["notes"])

    # formulas + styling rows 5-304
    for r in range(5, 305):
        jc = ws.cell(row=r, column=10)
        jc.value = '=IF($I{r}="","",IFERROR(VLOOKUP($I{r},Settings!$G$10:$H$15,2,FALSE),""))'.format(r=r)
        jc.number_format = PCT_FMT

        kc = ws.cell(row=r, column=11)
        kc.value = '=IF(OR($H{r}="",$J{r}=""),"",$H{r}*$J{r})'.format(r=r)
        kc.number_format = MONEY_FMT

        mc = ws.cell(row=r, column=13)
        mc.value = ('=IF(OR($I{r}="",$L{r}=""),"",IF(AND($L{r}<=TODAY(),'
                     '$I{r}<>Settings!$G$14,$I{r}<>Settings!$G$15),"CALL TODAY",""))').format(r=r)
        mc.font = Font(name="Arial", size=9, bold=True, color=RED)

        # input columns
        for col in (2, 3, 4, 5, 6, 7, 8, 9, 12, 14):
            cell = ws.cell(row=r, column=col)
            fmt = None
            if col in (2, 12):
                fmt = DATE_FMT
            elif col == 8:
                fmt = MONEY_FMT
            input_style(cell, number_format=fmt)

    dv_line = DataValidation(type="list", formula1="Settings!$C$10:$C$13", allow_blank=True)
    dv_source = DataValidation(type="list", formula1="Settings!$E$10:$E$15", allow_blank=True)
    dv_stage = DataValidation(type="list", formula1="Settings!$G$10:$G$15", allow_blank=True)
    ws.add_data_validation(dv_line)
    ws.add_data_validation(dv_source)
    ws.add_data_validation(dv_stage)
    dv_line.add("E5:E304")
    dv_source.add("F5:F304")
    dv_stage.add("I5:I304")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "B4:N304"

    widths = {"B": 11, "C": 24, "D": 14, "E": 13, "F": 14, "G": 32, "H": 14, "I": 12,
              "J": 12, "K": 13, "L": 13, "M": 12, "N": 28}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    return ws


# ---------------------------------------------------------------------------
# Jobs sheet
# ---------------------------------------------------------------------------

def build_jobs(wb, data):
    ws = wb.create_sheet("Jobs")
    apply_common(ws)
    build_nav_row(ws, "Jobs")
    build_banner(ws, "B2:L2", "  JOBS — won work only")
    build_ribbon(ws, "B3:L3", "  Job ref and margin fill themselves. Type the rest.")

    headers = ["Job ref", "Client", "Business line", "Value AED", "Est. cost AED", "Margin AED",
               "Margin %", "Status", "Start date", "Handover date", "Notes"]
    thin_navy = Side(style="thin", color=NAVY)
    for i, h in enumerate(headers):
        cell = ws.cell(row=4, column=2 + i)
        header_cell(cell, h)
        cell.border = Border(bottom=thin_navy)

    rows = data["jobs"]
    for i, rec in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=3, value=rec["client"])
        ws.cell(row=r, column=4, value=rec["line"])
        ws.cell(row=r, column=5, value=rec["value"])
        ws.cell(row=r, column=6, value=rec["est_cost"])
        ws.cell(row=r, column=9, value=rec["status"])
        ws.cell(row=r, column=10, value=parse_date(rec["start"]))
        ws.cell(row=r, column=11, value=parse_date(rec["handover"]))
        ws.cell(row=r, column=12, value=rec["notes"])

    for r in range(5, 205):
        bc = ws.cell(row=r, column=2)
        bc.value = '=IF($C{r}="","","RS-"&TEXT(ROW()-4,"000"))'.format(r=r)
        bc.font = Font(name="Arial", size=9, bold=True, color=NAVY)

        gc = ws.cell(row=r, column=7)
        gc.value = '=IF(OR($E{r}="",$F{r}=""),"",$E{r}-$F{r})'.format(r=r)
        gc.number_format = MONEY_FMT

        hc = ws.cell(row=r, column=8)
        hc.value = '=IF(OR($E{r}="",$E{r}=0,$G{r}=""),"",$G{r}/$E{r})'.format(r=r)
        hc.number_format = PCT_FMT

        for col in (3, 4, 5, 6, 9, 10, 11, 12):
            cell = ws.cell(row=r, column=col)
            fmt = None
            if col in (10, 11):
                fmt = DATE_FMT
            elif col in (5, 6):
                fmt = MONEY_FMT
            input_style(cell, number_format=fmt)

    dv_line = DataValidation(type="list", formula1="Settings!$C$10:$C$13", allow_blank=True)
    dv_status = DataValidation(type="list", formula1="Settings!$J$10:$J$11", allow_blank=True)
    ws.add_data_validation(dv_line)
    ws.add_data_validation(dv_status)
    dv_line.add("D5:D204")
    dv_status.add("I5:I204")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "B4:L204"

    widths = {"B": 9, "C": 24, "D": 13, "E": 13, "F": 13, "G": 12, "H": 10, "I": 12, "J": 12, "K": 14, "L": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    return ws


# ---------------------------------------------------------------------------
# Invoices sheet
# ---------------------------------------------------------------------------

def build_invoices(wb, data):
    ws = wb.create_sheet("Invoices")
    apply_common(ws)
    build_nav_row(ws, "Invoices")
    build_banner(ws, "B2:M2", "  INVOICES — who owes you money")
    build_ribbon(ws, "B3:M3", "  Balance, status and days overdue fill themselves. Overdue rows turn red, paid rows green.")

    headers = ["Invoice #", "Client", "Job ref", "Invoice date", "Amount AED", "Paid AED",
               "Paid date", "Balance AED", "Due date", "Status", "Days overdue", "Notes"]
    thin_navy = Side(style="thin", color=NAVY)
    for i, h in enumerate(headers):
        cell = ws.cell(row=4, column=2 + i)
        header_cell(cell, h)
        cell.border = Border(bottom=thin_navy)

    rows = data["invoices"]
    for i, rec in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=rec["inv"])
        ws.cell(row=r, column=3, value=rec["client"])
        ws.cell(row=r, column=4, value=rec["job_ref"])
        ws.cell(row=r, column=5, value=parse_date(rec["inv_date"]))
        ws.cell(row=r, column=6, value=rec["amount"])
        ws.cell(row=r, column=7, value=rec["paid"])
        ws.cell(row=r, column=8, value=parse_date(rec["paid_date"]))
        ws.cell(row=r, column=10, value=parse_date(rec["due"]))
        ws.cell(row=r, column=13, value=rec["notes"])

    for r in range(5, 205):
        ic = ws.cell(row=r, column=9)
        ic.value = '=IF($F{r}="","",$F{r}-N($G{r}))'.format(r=r)
        ic.number_format = MONEY_FMT

        kc = ws.cell(row=r, column=11)
        kc.value = ('=IF($F{r}="","",IF($I{r}<=0,"Paid",IF(AND($J{r}<>"",TODAY()>$J{r}),'
                     '"Overdue",IF(N($G{r})>0,"Partial","Upcoming"))))').format(r=r)

        lc = ws.cell(row=r, column=12)
        lc.value = ('=IF(OR($F{r}="",$I{r}<=0,$J{r}=""),"",IF(TODAY()>$J{r},TODAY()-$J{r},""))').format(r=r)
        lc.number_format = COUNT_FMT

        for col in (2, 3, 4, 5, 6, 7, 8, 10, 13):
            cell = ws.cell(row=r, column=col)
            fmt = None
            if col in (5, 8, 10):
                fmt = DATE_FMT
            elif col in (6, 7):
                fmt = MONEY_FMT
            input_style(cell, number_format=fmt)

    ws.conditional_formatting.add(
        "B5:M204",
        FormulaRule(formula=['$K5="Overdue"'], font=Font(color=RED, bold=True))
    )
    ws.conditional_formatting.add(
        "B5:M204",
        FormulaRule(formula=['$K5="Paid"'], font=Font(color=GREEN))
    )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "B4:M204"

    widths = {"B": 14, "C": 24, "D": 9, "E": 12, "F": 13, "G": 13, "H": 12, "I": 13,
              "J": 12, "K": 11, "L": 12, "M": 28}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    return ws


# ---------------------------------------------------------------------------
# Dashboard sheet
# ---------------------------------------------------------------------------

def build_dashboard(wb):
    ws = wb.create_sheet("Dashboard")
    apply_common(ws)
    build_nav_row(ws, "Dashboard")
    build_banner(ws, "B2:M2", "  DASHBOARD — read it, don't type in it")
    build_ribbon(ws, "B3:M3", "  Updates itself from Pipeline, Jobs and Invoices. Amounts in AED.")

    tile_border_side = Side(style="thin", color=TILE_BORDER_COLOR)
    tile_border = Border(left=tile_border_side, right=tile_border_side,
                          top=tile_border_side, bottom=tile_border_side)

    def tile(label_cell, value_cell, label_text, formula, number_format):
        lc = ws[label_cell]
        lc.value = label_text
        lc.font = Font(name="Arial", size=9, color=GRAY)
        lc.fill = FILL_CREAM
        lc.border = tile_border

        vc = ws[value_cell]
        vc.value = formula
        vc.font = Font(name="Arial", size=12, bold=True, color=NAVY)
        vc.fill = FILL_CREAM
        vc.border = tile_border
        vc.number_format = number_format

        row = int("".join(ch for ch in label_cell if ch.isdigit()))
        ws.row_dimensions[row].height = 20

    tile("B5", "C5", "Follow-ups due today / overdue",
         '=COUNTIF(Pipeline!$M$5:$M$304,"CALL TODAY")', COUNT_FMT)

    tile("B7", "C7", "Open pipeline (count)",
         '=COUNTIFS(Pipeline!$I$5:$I$304,"<>",Pipeline!$I$5:$I$304,"<>"&Settings!$G$14,'
         'Pipeline!$I$5:$I$304,"<>"&Settings!$G$15)', COUNT_FMT)
    tile("E7", "F7", "Open pipeline AED",
         '=SUMIFS(Pipeline!$H$5:$H$304,Pipeline!$I$5:$I$304,"<>",Pipeline!$I$5:$I$304,"<>"&Settings!$G$14,'
         'Pipeline!$I$5:$I$304,"<>"&Settings!$G$15)', MONEY_FMT)
    tile("H7", "I7", "Weighted pipeline AED",
         '=SUMIFS(Pipeline!$K$5:$K$304,Pipeline!$I$5:$I$304,"<>",Pipeline!$I$5:$I$304,"<>"&Settings!$G$14,'
         'Pipeline!$I$5:$I$304,"<>"&Settings!$G$15)', MONEY_FMT)

    tile("B9", "C9", "New leads this month",
         '=COUNTIFS(Pipeline!$B$5:$B$304,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
         'Pipeline!$B$5:$B$304,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', COUNT_FMT)
    tile("E9", "F9", "Jobs won this month (count)",
         '=COUNTIFS(Jobs!$J$5:$J$204,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
         'Jobs!$J$5:$J$204,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', COUNT_FMT)
    tile("H9", "I9", "Jobs won this month AED",
         '=SUMIFS(Jobs!$E$5:$E$204,Jobs!$J$5:$J$204,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
         'Jobs!$J$5:$J$204,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', MONEY_FMT)

    tile("B11", "C11", "Win rate (won ÷ closed)",
         '=IF(COUNTIF(Pipeline!$I$5:$I$304,Settings!$G$14)+COUNTIF(Pipeline!$I$5:$I$304,Settings!$G$15)=0,0,'
         'COUNTIF(Pipeline!$I$5:$I$304,Settings!$G$14)/(COUNTIF(Pipeline!$I$5:$I$304,Settings!$G$14)+'
         'COUNTIF(Pipeline!$I$5:$I$304,Settings!$G$15)))', PCT_FMT)
    tile("E11", "F11", "Invoiced this month AED",
         '=SUMIFS(Invoices!$F$5:$F$204,Invoices!$E$5:$E$204,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
         'Invoices!$E$5:$E$204,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', MONEY_FMT)
    tile("H11", "I11", "Collected this month AED",
         '=SUMIFS(Invoices!$G$5:$G$204,Invoices!$H$5:$H$204,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
         'Invoices!$H$5:$H$204,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))', MONEY_FMT)

    tile("B13", "C13", "Total outstanding AED", '=SUM(Invoices!$I$5:$I$204)', MONEY_FMT)
    tile("E13", "F13", "Overdue AED",
         '=SUMIFS(Invoices!$I$5:$I$204,Invoices!$K$5:$K$204,"Overdue")', MONEY_FMT)

    ws.conditional_formatting.add(
        "C5", FormulaRule(formula=["$C$5>0"], font=Font(color=RED, bold=True))
    )
    ws.conditional_formatting.add(
        "F13", FormulaRule(formula=["$F$13>0"], font=Font(color=RED, bold=True))
    )

    # Per-line table
    ws.merge_cells("B15:M15")
    bar = ws["B15"]
    bar.value = "  PIPELINE BY BUSINESS LINE"
    bar.font = Font(name="Arial", size=9, bold=True, color=NAVY)
    for row in ws["B15:M15"]:
        for c in row:
            c.fill = FILL_GOLD

    header_cell(ws["B16"], "Business line")
    header_cell(ws["C16"], "Open pipeline AED")
    header_cell(ws["D16"], "Won this year AED")
    header_cell(ws["E16"], "Open pipeline")

    for r in range(17, 21):
        s = r - 7
        bc = ws.cell(row=r, column=2)
        bc.value = '=IF(Settings!$C${s}="","",Settings!$C${s})'.format(s=s)
        bc.font = Font(name="Arial", size=10, bold=True, color=NAVY)

        cc = ws.cell(row=r, column=3)
        cc.value = ('=IF($B{r}="","",SUMIFS(Pipeline!$H$5:$H$304,Pipeline!$E$5:$E$304,$B{r},'
                     'Pipeline!$I$5:$I$304,"<>",Pipeline!$I$5:$I$304,"<>"&Settings!$G$14,'
                     'Pipeline!$I$5:$I$304,"<>"&Settings!$G$15))').format(r=r)
        cc.number_format = MONEY_FMT

        dc = ws.cell(row=r, column=4)
        dc.value = ('=IF($B{r}="","",SUMIFS(Jobs!$E$5:$E$204,Jobs!$D$5:$D$204,$B{r},'
                     'Jobs!$J$5:$J$204,">="&DATE(YEAR(TODAY()),1,1),'
                     'Jobs!$J$5:$J$204,"<"&DATE(YEAR(TODAY())+1,1,1)))').format(r=r)
        dc.number_format = MONEY_FMT

        ec = ws.cell(row=r, column=5)
        ec.value = ('=IF(OR($B{r}="",MAX($C$17:$C$20)=0),"",'
                     'REPT("█",ROUND($C{r}/MAX($C$17:$C$20)*20,0)))').format(r=r)
        ec.font = Font(name="Arial", size=10, color=GOLD)

    # Chart
    chart = BarChart()
    chart.type = "col"
    chart.title = "Pipeline value by business line"
    chart.width = 15
    chart.height = 7.5
    chart.legend = None
    data_ref = Reference(ws, min_col=3, min_row=17, max_row=20)
    cats_ref = Reference(ws, min_col=2, min_row=17, max_row=20)
    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cats_ref)
    chart.series[0].graphicalProperties.solidFill = "C9A227"
    ws.add_chart(chart, "G16")

    ws["B22"].value = ("Red tile means pick up the phone. The chart and this table follow whatever "
                        "names you give your business lines in Settings.")
    ws["B22"].font = Font(name="Arial", size=8, color=GRAY)

    widths = {"B": 26, "C": 15, "D": 2, "E": 26, "F": 15, "G": 2, "H": 26, "I": 15,
              "J": 10, "K": 10, "L": 10, "M": 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    return ws


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    data = load_data()

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    build_start_here(wb)
    build_dashboard(wb)
    build_pipeline(wb, data)
    build_jobs(wb, data)
    build_invoices(wb, data)
    build_settings(wb, data)

    wb.save(str(OUT_FILE))

    n_pipeline = len(data["pipeline"])
    n_jobs = len(data["jobs"])
    n_invoices = len(data["invoices"])
    print("Wrote {} with sheets {} | pipeline={} jobs={} invoices={}".format(
        OUT_FILE.name, SHEET_ORDER, n_pipeline, n_jobs, n_invoices))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        import traceback
        traceback.print_exc()
        sys.exit(1)
