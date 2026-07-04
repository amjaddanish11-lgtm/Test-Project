#!/usr/bin/env python3
"""Independent verifier for RS_Command_Center.xlsx.

Written separately from build_cockpit.py, on purpose. It:
  1. Recomputes every Dashboard KPI in pure Python from the dataset in
     rs_checks.json (TODAY()-dependent values are recomputed with
     date.today(), matching what Excel formulas produce at recalc time).
  2. Reloads the workbook and checks structure: sheet order, gridlines,
     freeze panes, auto-filters, protection, dropdown sources, number
     formats, merged banners, locked/unlocked cells.
  3. Parses the Dashboard/data-tab formula strings and label-checks every
     Settings reference (e.g. the cell a formula uses as "Won" must
     actually say Won).
  4. If the file has been opened and SAVED by real Excel, compares the
     cached formula results against the Python-recomputed expectations.
     Before that first Excel save there are no cached values; that step
     is skipped with a notice (openpyxl cannot recalculate formulas).

Exit code 0 = all checks passed (skips are allowed), 1 = any failure.
"""
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
XLSX = HERE / "RS_Command_Center.xlsx"
CHECKS = HERE / "rs_checks.json"

failures = []
passes = 0
skips = []


def ok(label):
    global passes
    passes += 1


def fail(label, detail):
    failures.append(f"FAIL: {label}: {detail}")


def check(label, cond, detail=""):
    if cond:
        ok(label)
    else:
        fail(label, detail)


# ---------------------------------------------------------------- expectations
def month_bounds(today):
    start = today.replace(day=1)
    nxt = (start + timedelta(days=32)).replace(day=1)
    return start, nxt


def compute_expected(ds, today):
    """Pure-Python recomputation of every KPI from the raw dataset."""
    P = {s["stage"]: s for s in ds["settings"]["stages"]}
    pipe = ds["pipeline"]
    jobs = ds["jobs"]
    inv = ds["invoices"]
    d = date.fromisoformat
    mstart, mnext = month_bounds(today)

    opn = [r for r in pipe if not P[r["stage"]]["closed"]]
    won = [r for r in pipe if r["stage"] == "Won"]
    lost = [r for r in pipe if r["stage"] == "Lost"]
    closed = len(won) + len(lost)

    exp = {
        "followups_due_today_or_overdue": sum(
            1 for r in opn if r["next_followup"] and d(r["next_followup"]) <= today
        ),
        "open_pipeline_count": len(opn),
        "open_pipeline_total_aed": sum(r["value"] for r in opn),
        "open_pipeline_weighted_aed": sum(
            r["value"] * P[r["stage"]]["probability"] for r in opn
        ),
        "new_leads_this_month": sum(
            1 for r in pipe if mstart <= d(r["date_in"]) < mnext
        ),
        "jobs_won_this_month_count": sum(
            1 for j in jobs if mstart <= d(j["start"]) < mnext
        ),
        "jobs_won_this_month_aed": sum(
            j["value"] for j in jobs if mstart <= d(j["start"]) < mnext
        ),
        "win_rate": (len(won) / closed) if closed else 0,
        "invoiced_this_month_aed": sum(
            i["amount"] for i in inv if mstart <= d(i["inv_date"]) < mnext
        ),
        "collected_this_month_aed": sum(
            i["paid"]
            for i in inv
            if i["paid_date"] and mstart <= d(i["paid_date"]) < mnext
        ),
        "total_outstanding_aed": sum(i["amount"] - i["paid"] for i in inv),
        "overdue_aed": sum(
            i["amount"] - i["paid"]
            for i in inv
            if i["amount"] - i["paid"] > 0 and d(i["due"]) < today
        ),
    }
    lines = ds["settings"]["business_lines"]
    exp["per_line_open_aed"] = {
        L: sum(r["value"] for r in opn if r["line"] == L) for L in lines
    }
    exp["per_line_won_year_aed"] = {
        L: sum(
            j["value"]
            for j in jobs
            if j["line"] == L and d(j["start"]).year == today.year
        )
        for L in lines
    }
    exp["invoice_status"] = {}
    for i in inv:
        bal = i["amount"] - i["paid"]
        if bal <= 0:
            st, days = "Paid", None
        elif d(i["due"]) < today:
            st, days = "Overdue", (today - d(i["due"])).days
        elif i["paid"] > 0:
            st, days = "Partial", None
        else:
            st, days = "Upcoming", None
        exp["invoice_status"][i["inv"]] = {
            "status": st, "balance": bal, "days_overdue": days
        }
    exp["call_today_clients"] = sorted(
        r["client"]
        for r in opn
        if r["next_followup"] and d(r["next_followup"]) <= today
    )
    exp["job_margins"] = {}
    for n, j in enumerate(jobs):
        ref = f"RS-{n + 1:03d}"
        exp["job_margins"][ref] = {
            "margin": j["value"] - j["est_cost"],
            "margin_pct": (j["value"] - j["est_cost"]) / j["value"],
        }
    return exp


# ------------------------------------------------------- static cross-check
def crosscheck_static(exp, checks):
    """The frozen hand-computed values in rs_checks.json must agree with this
    script's independent recomputation for all non-dynamic checks."""
    static_map = {
        "open_pipeline_count": exp["open_pipeline_count"],
        "open_pipeline_total_aed": exp["open_pipeline_total_aed"],
        "open_pipeline_weighted_aed": exp["open_pipeline_weighted_aed"],
        "win_rate": exp["win_rate"],
        "total_outstanding_aed": exp["total_outstanding_aed"],
    }
    for key, got in static_map.items():
        want = checks[key]["value"]
        check(f"static:{key}", abs(got - want) < 1e-9, f"json={want} recomputed={got}")
    for L, want in checks["per_line_open_pipeline_aed"]["value"].items():
        got = exp["per_line_open_aed"][L]
        check(f"static:per_line_open:{L}", got == want, f"json={want} recomputed={got}")
    for ref, m in checks["job_margins"]["value"].items():
        got = exp["job_margins"][ref]
        check(f"static:margin:{ref}", got["margin"] == m["margin"], f"{got} vs {m}")
        check(
            f"static:margin_pct:{ref}",
            abs(got["margin_pct"] - m["margin_pct"]) <= checks["job_margins"]["margin_pct_tolerance"],
            f"{got['margin_pct']} vs {m['margin_pct']}",
        )


# ------------------------------------------------------------------ structure
SHEETS = ["Start Here", "Dashboard", "Pipeline", "Jobs", "Invoices", "Settings"]

SETTINGS_LABELS = {
    "C10": "Signage", "C11": "Neon", "C12": "Fitout", "C13": "Website",
    "E10": "Google Ads", "E15": "Designer partner",
    "G10": "New", "G13": "Negotiating", "G14": "Won", "G15": "Lost",
    "J10": "In progress", "J11": "Done",
}
SETTINGS_PROBS = {"H10": 0.10, "H11": 0.25, "H12": 0.40, "H13": 0.60,
                  "H14": 1.00, "H15": 0.00}

# Formula fragments each Dashboard KPI cell must contain (range + Settings refs)
KPI_FORMULA_MUST_CONTAIN = {
    "followups_due_today_or_overdue": ["Pipeline!$M$5:$M$304", "CALL TODAY"],
    "open_pipeline_count": ["Pipeline!$I$5:$I$304", "Settings!$G$14", "Settings!$G$15"],
    "open_pipeline_total_aed": ["Pipeline!$H$5:$H$304", "Pipeline!$I$5:$I$304",
                                "Settings!$G$14", "Settings!$G$15"],
    "open_pipeline_weighted_aed": ["Pipeline!$K$5:$K$304", "Settings!$G$14",
                                   "Settings!$G$15"],
    "new_leads_this_month": ["Pipeline!$B$5:$B$304", "TODAY()"],
    "jobs_won_this_month_count": ["Jobs!$J$5:$J$204", "TODAY()"],
    "jobs_won_this_month_aed": ["Jobs!$E$5:$E$204", "Jobs!$J$5:$J$204", "TODAY()"],
    "win_rate": ["Settings!$G$14", "Settings!$G$15", "=0,0,"],
    "invoiced_this_month_aed": ["Invoices!$F$5:$F$204", "Invoices!$E$5:$E$204"],
    "collected_this_month_aed": ["Invoices!$G$5:$G$204", "Invoices!$H$5:$H$204"],
    "total_outstanding_aed": ["Invoices!$I$5:$I$204"],
    "overdue_aed": ["Invoices!$I$5:$I$204", "Invoices!$K$5:$K$204", "Overdue"],
}


def check_workbook(cj, exp):
    import openpyxl

    wb = openpyxl.load_workbook(XLSX)
    check("sheet order", wb.sheetnames == SHEETS, f"got {wb.sheetnames}")

    for sn in wb.sheetnames:
        ws = wb[sn]
        check(f"{sn}: gridlines off", ws.sheet_view.showGridLines is False,
              str(ws.sheet_view.showGridLines))
        check(f"{sn}: protected, no password",
              ws.protection.sheet and not ws.protection.password,
              f"sheet={ws.protection.sheet} pw={bool(ws.protection.password)}")
        check(f"{sn}: autoFilter allowed under protection",
              ws.protection.autoFilter is False, str(ws.protection.autoFilter))
        wA = ws.column_dimensions["A"].width
        check(f"{sn}: column A gutter", wA is not None and wA <= 4, f"width={wA}")

    for sn, flt in [("Pipeline", "B4:N304"), ("Jobs", "B4:L204"),
                    ("Invoices", "B4:M204")]:
        ws = wb[sn]
        check(f"{sn}: freeze A5", ws.freeze_panes == "A5", str(ws.freeze_panes))
        check(f"{sn}: autofilter", ws.auto_filter.ref == flt, str(ws.auto_filter.ref))

    # banners merged full width on every sheet
    for sn in SHEETS:
        merges = [str(r) for r in wb[sn].merged_cells.ranges]
        check(f"{sn}: row2 banner merged", any(m.startswith("B2:") for m in merges),
              f"merges={merges}")

    # Settings labels + probabilities (the label-check that guards every
    # Settings!$G$14 / $G$15 style reference used by formulas)
    st = wb["Settings"]
    for addr, want in SETTINGS_LABELS.items():
        check(f"Settings {addr}='{want}'", st[addr].value == want,
              f"got {st[addr].value!r}")
    for addr, want in SETTINGS_PROBS.items():
        got = st[addr].value
        check(f"Settings {addr}={want}", isinstance(got, (int, float))
              and abs(got - want) < 1e-9, f"got {got!r}")

    # Dashboard KPI formulas reference the intended ranges
    dash = wb["Dashboard"]
    cmap = cj["dashboard_cell_map"]
    for kpi, frags in KPI_FORMULA_MUST_CONTAIN.items():
        cell = cmap[kpi]
        f = dash[cell].value
        if not isinstance(f, str) or not f.startswith("="):
            fail(f"dashboard {kpi} @{cell}", f"no formula, got {f!r}")
            continue
        for frag in frags:
            check(f"dashboard {kpi} @{cell} contains {frag}", frag in f,
                  f"formula={f}")

    # per-line table: line cells reference Settings C10:C13; SUMIFS wired
    tbl = cmap["per_line_table"]
    for i, (line, row) in enumerate(tbl["rows"].items()):
        b = dash[f"{tbl['line_col']}{row}"].value
        check(f"dashboard line row {row} -> Settings!$C${10+i}",
              isinstance(b, str) and f"Settings!$C${10+i}" in b, f"got {b!r}")
        c = dash[f"{tbl['open_aed_col']}{row}"].value
        check(f"dashboard line {line} open SUMIFS",
              isinstance(c, str) and "Pipeline!$H$5:$H$304" in c
              and "Pipeline!$E$5:$E$304" in c, f"got {c!r}")
        dcol = dash[f"{tbl['won_year_col']}{row}"].value
        check(f"dashboard line {line} won-year SUMIFS",
              isinstance(dcol, str) and "Jobs!$E$5:$E$204" in dcol
              and "Jobs!$D$5:$D$204" in dcol and "Jobs!$J$5:$J$204" in dcol,
              f"got {dcol!r}")
        e = dash[f"{tbl['rept_bar_col']}{row}"].value
        check(f"dashboard line {line} REPT bar",
              isinstance(e, str) and "REPT(" in e and "MAX($C$17:$C$20)" in e,
              f"got {e!r}")

    check("dashboard chart present", len(dash._charts) == 1,
          f"{len(dash._charts)} charts")

    # data-tab formulas, dropdowns, guards
    pl = wb["Pipeline"]
    check("Pipeline J5 stage lookup",
          pl["J5"].value == '=IF($I5="","",IFERROR(VLOOKUP($I5,Settings!$G$10:$H$15,2,FALSE),""))',
          repr(pl["J5"].value))
    check("Pipeline K5 weighted", pl["K5"].value == '=IF(OR($H5="",$J5=""),"",$H5*$J5)',
          repr(pl["K5"].value))
    m5 = pl["M5"].value
    check("Pipeline M5 flag", isinstance(m5, str) and "CALL TODAY" in m5
          and "$L5<=TODAY()" in m5 and "Settings!$G$14" in m5
          and "Settings!$G$15" in m5, repr(m5))
    check("Pipeline M304 flag filled to row 304",
          isinstance(pl["M304"].value, str) and "CALL TODAY" in pl["M304"].value,
          repr(pl["M304"].value))

    jb = wb["Jobs"]
    check("Jobs B5 auto ref",
          jb["B5"].value == '=IF($C5="","","RS-"&TEXT(ROW()-4,"000"))',
          repr(jb["B5"].value))
    check("Jobs H5 margin pct guarded",
          jb["H5"].value == '=IF(OR($E5="",$E5=0,$G5=""),"",$G5/$E5)',
          repr(jb["H5"].value))

    iv = wb["Invoices"]
    k5 = iv["K5"].value
    check("Invoices K5 status order (Paid, then Overdue beats Partial)",
          isinstance(k5, str)
          and k5.index("Paid") < k5.index("Overdue") < k5.index("Partial")
          and "Upcoming" in k5, repr(k5))
    check("Invoices I5 balance", iv["I5"].value == '=IF($F5="","",$F5-N($G5))',
          repr(iv["I5"].value))
    l5 = iv["L5"].value
    check("Invoices L5 days overdue blank-guarded",
          isinstance(l5, str) and "$I5<=0" in l5 and "TODAY()-$J5" in l5, repr(l5))

    dv_expect = {
        "Pipeline": {"Settings!$C$10:$C$13", "Settings!$E$10:$E$15",
                     "Settings!$G$10:$G$15"},
        "Jobs": {"Settings!$C$10:$C$13", "Settings!$J$10:$J$11"},
    }
    for sn, want in dv_expect.items():
        got = {dv.formula1 for dv in wb[sn].data_validations.dataValidation}
        check(f"{sn}: dropdown sources", want <= got, f"got {got}")
        bad = [f for f in got if f.startswith("=")]
        check(f"{sn}: no '=' in formula1", not bad, str(bad))

    # locked vs unlocked spot checks
    for sn, unlocked, locked in [
        ("Pipeline", ["B5", "I5", "L5", "N304"], ["J5", "K5", "M5"]),
        ("Jobs", ["C5", "J5", "L204"], ["B5", "G5", "H5"]),
        ("Invoices", ["B5", "E5", "G5", "J5"], ["I5", "K5", "L5"]),
        ("Settings", ["C5", "C10", "E10", "G10", "H15", "J11"], []),
        ("Dashboard", [], ["C5", "F13", "C17"]),
    ]:
        ws = wb[sn]
        for a in unlocked:
            check(f"{sn} {a} unlocked", ws[a].protection.locked is False, "locked")
        for a in locked:
            check(f"{sn} {a} locked", ws[a].protection.locked is not False, "unlocked")

    # number format spot checks
    fmt = [
        ("Pipeline", "B5", "DD/MM/YYYY"), ("Pipeline", "H5", "#,##0.00"),
        ("Pipeline", "J5", "0.0%"), ("Pipeline", "L5", "DD/MM/YYYY"),
        ("Jobs", "E5", "#,##0.00"), ("Jobs", "H5", "0.0%"),
        ("Jobs", "J5", "DD/MM/YYYY"),
        ("Invoices", "E5", "DD/MM/YYYY"), ("Invoices", "F5", "#,##0.00"),
        ("Invoices", "J5", "DD/MM/YYYY"),
        ("Dashboard", "F7", "#,##0.00"), ("Dashboard", "C11", "0.0%"),
    ]
    for sn, a, want in fmt:
        got = wb[sn][a].number_format
        check(f"{sn} {a} format {want}", got.upper() == want.upper(), f"got {got}")

    # dummy rows landed intact (values match dataset)
    ds = cj["dataset"]
    from datetime import datetime as _dt
    def as_date(v):
        return v.date() if isinstance(v, _dt) else v
    for i, r in enumerate(ds["pipeline"]):
        row = 5 + i
        check(f"Pipeline row {row} client", pl[f"C{row}"].value == r["client"],
              f"got {pl[f'C{row}'].value!r} want {r['client']!r}")
        check(f"Pipeline row {row} value", pl[f"H{row}"].value == r["value"],
              f"got {pl[f'H{row}'].value!r}")
        check(f"Pipeline row {row} stage", pl[f"I{row}"].value == r["stage"],
              f"got {pl[f'I{row}'].value!r}")
        want_fu = date.fromisoformat(r["next_followup"]) if r["next_followup"] else None
        check(f"Pipeline row {row} follow-up", as_date(pl[f"L{row}"].value) == want_fu,
              f"got {pl[f'L{row}'].value!r} want {want_fu}")
    for i, j in enumerate(ds["jobs"]):
        row = 5 + i
        check(f"Jobs row {row} value/cost",
              jb[f"E{row}"].value == j["value"] and jb[f"F{row}"].value == j["est_cost"],
              f"got {jb[f'E{row}'].value},{jb[f'F{row}'].value}")
    for i, v in enumerate(ds["invoices"]):
        row = 5 + i
        check(f"Invoices row {row} amount/paid",
              iv[f"F{row}"].value == v["amount"] and (iv[f"G{row}"].value or 0) == v["paid"],
              f"got {iv[f'F{row}'].value},{iv[f'G{row}'].value}")
        check(f"Invoices row {row} due", as_date(iv[f"J{row}"].value) == date.fromisoformat(v["due"]),
              f"got {iv[f'J{row}'].value!r}")

    # Arial everywhere - spot-check computed/formula cells (historical defect:
    # cells given only a value + number_format fall back to Calibri)
    for sn, cells in [
        ("Pipeline", ["J5", "K5", "M5", "K304"]),
        ("Jobs", ["B5", "G5", "H5", "H204"]),
        ("Invoices", ["I5", "K5", "L5", "L204"]),
        ("Dashboard", ["C5", "C17", "D17", "E17"]),
        ("Settings", ["C10", "G14"]),
        ("Start Here", ["B5", "B7"]),
    ]:
        for a in cells:
            got = wb[sn][a].font.name
            check(f"{sn} {a} font Arial", got == "Arial", f"got {got!r}")

    # conditional formatting present
    dcf = {str(rng): [r.type for r in rules]
           for rng, rules in dash.conditional_formatting._cf_rules.items()}
    icf = list(iv.conditional_formatting._cf_rules.items())
    check("Dashboard CF on C5 and F13",
          any("C5" in str(r.sqref) for r, _ in dash.conditional_formatting._cf_rules.items())
          and any("F13" in str(r.sqref) for r, _ in dash.conditional_formatting._cf_rules.items()),
          str(dcf))
    check("Invoices CF rules (2 over table)",
          sum(len(rules) for _, rules in icf) >= 2
          and any("B5:M204" in str(r.sqref) for r, _ in icf), str(icf))


# ------------------------------------------------- cached values (post-Excel)
def check_cached_values(cj, exp):
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    dash = wb["Dashboard"]
    cmap = cj["dashboard_cell_map"]
    probe = dash[cmap["open_pipeline_count"]].value
    if probe is None:
        skips.append(
            "Cached formula values not present - the file has not yet been "
            "opened, recalculated and saved by real Excel. Run this script "
            "again after Excel saves it to compare Dashboard values."
        )
        return
    numeric = {
        "followups_due_today_or_overdue": 0,
        "open_pipeline_count": 0,
        "open_pipeline_total_aed": 0.01,
        "open_pipeline_weighted_aed": 0.01,
        "new_leads_this_month": 0,
        "jobs_won_this_month_count": 0,
        "jobs_won_this_month_aed": 0.01,
        "win_rate": 0.0001,
        "invoiced_this_month_aed": 0.01,
        "collected_this_month_aed": 0.01,
        "total_outstanding_aed": 0.01,
        "overdue_aed": 0.01,
    }
    for kpi, tol in numeric.items():
        got = dash[cmap[kpi]].value
        want = exp[kpi]
        good = isinstance(got, (int, float)) and abs(got - want) <= tol
        check(f"recalc:{kpi}", good, f"cell={got!r} expected={want}")
    tbl = cmap["per_line_table"]
    for line, row in tbl["rows"].items():
        got_open = dash[f"{tbl['open_aed_col']}{row}"].value
        got_won = dash[f"{tbl['won_year_col']}{row}"].value
        check(f"recalc:line {line} open", got_open == exp["per_line_open_aed"][line],
              f"cell={got_open!r} expected={exp['per_line_open_aed'][line]}")
        check(f"recalc:line {line} won-yr", got_won == exp["per_line_won_year_aed"][line],
              f"cell={got_won!r} expected={exp['per_line_won_year_aed'][line]}")
    iv = wb["Invoices"]
    for i, (inv_id, st) in enumerate(exp["invoice_status"].items()):
        row = 5 + i
        check(f"recalc:{inv_id} status", iv[f"K{row}"].value == st["status"],
              f"cell={iv[f'K{row}'].value!r} expected={st['status']}")
        check(f"recalc:{inv_id} balance", iv[f"I{row}"].value == st["balance"],
              f"cell={iv[f'I{row}'].value!r} expected={st['balance']}")


def main():
    cj = json.loads(CHECKS.read_text())
    today = date.today()
    exp = compute_expected(cj["dataset"], today)
    crosscheck_static(exp, cj["checks"])
    if not XLSX.exists():
        fail("workbook", f"{XLSX} not found - run build_cockpit.py first")
    else:
        check_workbook(cj, exp)
        check_cached_values(cj, exp)

    print(f"post_check: {passes} passed, {len(failures)} failed, "
          f"{len(skips)} skipped (today={today})")
    for s in skips:
        print("  SKIP:", s)
    for f in failures:
        print(" ", f)
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
